import csv
import io
import zipfile

import pytest

from cn_extras.extract import extract
from cn_extras.extract_local import ExtractionFailed
from cn_extras.office import pptx_text, xlsx_text
from tests.cn_extras import fixtures as fx


def _sheet_csv(text: str, heading: str) -> list:
    block = text.split(f"=== {heading} ===\n", 1)[1].split("\n\n=== ", 1)[0]
    return list(csv.reader(io.StringIO(block)))


class TestXlsx:
    def test_lists_every_sheet_with_hidden_flag(self):
        book = xlsx_text(fx.real_xlsx())
        assert book.sheet_names == ["Inscriptions", "Budget", "Interne", "Vide"]
        first_line = book.text.splitlines()[0]
        assert "1. Inscriptions" in first_line and "3. Interne (hidden)" in first_line

    def test_cells_are_valid_csv_with_typed_values(self):
        rows = _sheet_csv(xlsx_text(fx.real_xlsx()).text, "Sheet 1: Inscriptions")
        assert rows[0] == ["Nom", "Pays", "Âge", "Arrivée", "Payé"]
        assert rows[1] == ["João", "Brésil", "23", "2027-07-26", "TRUE"]
        assert rows[2] == [
            'Marie, dite "Mimi"',
            "France",
            "19.5",
            "2027-07-27 14:30:00",
            "FALSE",
        ]

    def test_float_integers_print_without_decimal(self):
        rows = _sheet_csv(xlsx_text(fx.real_xlsx()).text, "Sheet 2: Budget")
        assert rows[1] == ["Transport", "1200"]

    def test_empty_sheet_is_marked(self):
        assert "=== Sheet 4: Vide ===\n(empty sheet)" in xlsx_text(fx.real_xlsx()).text

    @pytest.mark.parametrize("selector", ["2", "budget", " Budget "])
    def test_select_one_sheet_by_number_or_name(self, selector):
        text = xlsx_text(fx.real_xlsx(), selector).text
        assert "=== Sheet 2: Budget ===" in text
        assert "=== Sheet 1" not in text

    def test_unknown_sheet_lists_choices(self):
        with pytest.raises(ExtractionFailed, match="1. Inscriptions, 2. Budget"):
            xlsx_text(fx.real_xlsx(), "Nope")

    def test_output_cap(self, monkeypatch):
        import cn_extras.office as office

        monkeypatch.setattr(office, "MAX_OUTPUT_CHARS", 50)
        book = xlsx_text(fx.real_xlsx())
        assert book.notes and "capped" in book.notes[0]
        assert "=== Sheet 2" not in book.text

    def test_zip_bomb_is_refused_before_parsing(self, monkeypatch):
        monkeypatch.setenv("CN_OFFICE_MAX_UNCOMPRESSED_BYTES", "1000")
        with pytest.raises(ExtractionFailed, match="uncompressed"):
            xlsx_text(fx.real_xlsx())

    def test_damaged(self):
        with pytest.raises(ExtractionFailed):
            xlsx_text(b"not a zip")

    def test_via_extract_with_sheet(self):
        r = extract(fx.real_xlsx(), fx.XLSX, "b.xlsx", sheet="Budget")
        assert r.kind == "text" and "Transport" in r.text and "João" not in r.text

    def test_unknown_sheet_via_extract_is_a_clear_reason(self):
        r = extract(fx.real_xlsx(), fx.XLSX, "b.xlsx", sheet="9")
        assert r.kind == "unsupported" and "Inscriptions" in r.reason


class TestPptx:
    def test_slides_in_numeric_order_with_notes_and_no_field_text(self):
        text = pptx_text(fx.pptx_with_notes())
        assert (
            text.index("slide 1 ---")
            < text.index("slide 2 ---")
            < text.index("slide 10 ---")
        )
        assert "Speaker notes: Dire bonjour aux pèlerins" in text
        assert "99" not in text

    def test_notes_can_be_left_out(self):
        assert "Speaker notes" not in pptx_text(
            fx.pptx_with_notes(), include_notes=False
        )

    def test_no_slides(self):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr("ppt/presentation.xml", "<x/>")
        with pytest.raises(ExtractionFailed, match="no slides"):
            pptx_text(buf.getvalue())

    def test_entity_bomb_is_refused(self):
        bomb = (
            '<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "aaaaaaaaaa">'
            '<!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;&a;&a;">]><x>&b;</x>'
        )
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr("ppt/slides/slide1.xml", bomb)
        with pytest.raises(ExtractionFailed):
            pptx_text(buf.getvalue())


class TestHostileWorkbooks:
    """Phase 2 review: sparse/wide sheets and oversized parts."""

    @staticmethod
    def _save(wb) -> bytes:
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def test_sparse_sheet_with_huge_dimension_is_fast(self):
        import time

        from openpyxl import Workbook

        wb = Workbook()
        wb.active["A1"] = "top"
        wb.active["XFD1048576"] = "far"
        started = time.monotonic()
        text = xlsx_text(self._save(wb)).text
        assert time.monotonic() - started < 5
        assert "top" in text and "far" in text

    def test_wide_rows_hit_the_cell_budget(self, monkeypatch):
        import cn_extras.office as office
        from openpyxl import Workbook

        monkeypatch.setattr(office, "MAX_CELLS_VISITED", 100_000)
        wb = Workbook()
        for r in range(1, 50):
            wb.active.cell(row=r, column=1, value="a")
            wb.active.cell(row=r, column=16384, value="z")
        book = xlsx_text(self._save(wb))
        assert any("stopped after 100000 cells" in n for n in book.notes)

    @staticmethod
    def _zip(members) -> bytes:
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
            for name, size in members.items():
                zf.writestr(name, "x" * size)
        return buf.getvalue()

    def test_parsed_part_over_per_part_limit_is_refused(self, monkeypatch):
        from cn_extras.office import _open_zip

        monkeypatch.setenv("WORKSPACE_MCP_MAX_OFFICE_XML_BYTES", "10000")
        with pytest.raises(ExtractionFailed, match="per-part"):
            _open_zip(self._zip({"xl/sharedStrings.xml": 50_000}))
        with pytest.raises(ExtractionFailed, match="per-part"):
            _open_zip(self._zip({"ppt/slides/slide1.xml": 50_000}))

    def test_streamed_worksheet_may_exceed_per_part_limit(self, monkeypatch):
        from cn_extras.office import _open_zip

        monkeypatch.setenv("WORKSPACE_MCP_MAX_OFFICE_XML_BYTES", "10000")
        _open_zip(
            self._zip({"xl/worksheets/sheet1.xml": 50_000, "xl/workbook.xml": 100})
        ).close()
