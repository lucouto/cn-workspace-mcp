"""Structure-aware readers for XLSX (per sheet, as CSV) and PPTX (per slide, with notes).

Upstream's ``extract_office_xml_text`` flattens a workbook into one line of
cell values with no sheet names, and reads PPTX slides without speaker notes.
Both matter for Drive exports of Sheets and Slides, so cn_extras reads these
two formats itself. DOCX still goes through upstream.

Both readers are bounded: the total uncompressed size of the ZIP is checked
before anything is parsed (zip bombs), and the produced text is capped.
"""

import csv
import datetime as dt
import io
import os
import posixpath
import re
import zipfile
from dataclasses import dataclass, field
from typing import List, Optional

from defusedxml import DefusedXmlException
from defusedxml import ElementTree as ET

from cn_extras.extract_local import ExtractionFailed

# Whole-package ceiling. Parts parsed into memory (sharedStrings, styles, each
# slide) are held to upstream's per-part WORKSPACE_MCP_MAX_OFFICE_XML_BYTES
# (default 25 MiB), since a parsed tree costs 14-30x the XML size; worksheets
# are streamed by openpyxl and only count toward this total.
DEFAULT_MAX_UNCOMPRESSED_BYTES = 256 * 1024 * 1024
# Far more than one read returns (pagination), small enough to bound memory.
MAX_OUTPUT_CHARS = 5_000_000
# Cells visited per workbook, empty ones included. openpyxl pads rows to the
# sheet width, so a tiny file with cells in A and XFD would otherwise cost
# 16384 cells per row.
MAX_CELLS_VISITED = 5_000_000
_STREAMED_PARTS = re.compile(r"^xl/worksheets/sheet\d+\.xml$")

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_REL = "{http://schemas.openxmlformats.org/package/2006/relationships}"


def _max_uncompressed_bytes() -> int:
    raw = os.getenv("CN_OFFICE_MAX_UNCOMPRESSED_BYTES", "").strip()
    if not raw:
        return DEFAULT_MAX_UNCOMPRESSED_BYTES
    try:
        value = int(raw)
    except ValueError as exc:
        raise ValueError(
            f"Invalid CN_OFFICE_MAX_UNCOMPRESSED_BYTES={raw!r}; expected a positive integer."
        ) from exc
    if value <= 0:
        raise ValueError(
            f"Invalid CN_OFFICE_MAX_UNCOMPRESSED_BYTES={raw!r}; expected a positive integer."
        )
    return value


def _open_zip(data: bytes) -> zipfile.ZipFile:
    from core.file_limits import get_max_office_xml_bytes

    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise ExtractionFailed("the file is not a valid Office document.") from exc
    # zipfile never yields more than an entry's declared file_size, so the
    # declared sizes bound what parsing can expand to.
    mb = 1024 * 1024
    total = sum(info.file_size for info in zf.infolist())
    limit = _max_uncompressed_bytes()
    per_part = get_max_office_xml_bytes()
    problem = None
    if total > limit:
        problem = (
            f"expands to {total // mb} MB uncompressed, over the {limit // mb} MB limit"
        )
    elif per_part:
        for info in zf.infolist():
            if info.file_size > per_part and not _STREAMED_PARTS.match(info.filename):
                problem = (
                    f"has a {info.file_size // mb} MB internal part, over the "
                    f"{per_part // mb} MB per-part limit"
                )
                break
    if problem:
        zf.close()
        raise ExtractionFailed(f"the document {problem}.")
    return zf


# --- XLSX ----------------------------------------------------------------------


@dataclass
class WorkbookText:
    text: str
    sheet_names: List[str]
    notes: List[str] = field(default_factory=list)


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, dt.datetime):
        return (
            value.date().isoformat()
            if value.time() == dt.time()
            else value.isoformat(sep=" ")
        )
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    return str(value)


def _select_sheets(worksheets, sheet: Optional[str]):
    if sheet is None or str(sheet).strip().lower() in ("", "all"):
        return list(enumerate(worksheets, 1))
    wanted = str(sheet).strip()
    if wanted.isdigit():
        index = int(wanted)
        if 1 <= index <= len(worksheets):
            return [(index, worksheets[index - 1])]
    for i, ws in enumerate(worksheets, 1):
        if ws.title.lower() == wanted.lower():
            return [(i, ws)]
    names = ", ".join(f"{i}. {ws.title}" for i, ws in enumerate(worksheets, 1))
    raise ExtractionFailed(
        f"no sheet '{wanted}'. Sheets (use the number or the name): {names}"
    )


def xlsx_text(data: bytes, sheet: Optional[str] = None) -> WorkbookText:
    """Return each selected sheet as CSV under a ``=== Sheet N: name ===`` heading."""
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    _open_zip(data).close()
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (InvalidFileException, KeyError, zipfile.BadZipFile, ValueError) as exc:
        raise ExtractionFailed("the spreadsheet appears damaged.") from exc

    try:
        worksheets = wb.worksheets
        names = [ws.title for ws in worksheets]
        chosen = _select_sheets(worksheets, sheet)
        listing = ", ".join(
            f"{i}. {ws.title}" + (" (hidden)" if ws.sheet_state != "visible" else "")
            for i, ws in enumerate(worksheets, 1)
        )
        out = io.StringIO()
        out.write(f"Workbook with {len(worksheets)} sheet(s): {listing}\n")
        notes: List[str] = []
        visited = 0
        for index, ws in chosen:
            out.write(f"\n=== Sheet {index}: {ws.title} ===\n")
            writer = csv.writer(out, lineterminator="\n")
            if hasattr(ws, "reset_dimensions"):
                # Ignore the declared <dimension>: otherwise every row is padded
                # to max_column and every gap up to max_row is materialised.
                ws.reset_dimensions()
            rows = 0
            cap = None
            for row in ws.iter_rows(values_only=True):
                visited += len(row) + 1
                if visited > MAX_CELLS_VISITED:
                    cap = f"stopped after {MAX_CELLS_VISITED} cells (sparse or very wide sheet)"
                    break
                last = len(row)
                while last and row[last - 1] in (None, ""):
                    last -= 1
                if not last:
                    continue
                writer.writerow([_cell(v) for v in row[:last]])
                rows += 1
                if out.tell() > MAX_OUTPUT_CHARS:
                    cap = f"output capped at {MAX_OUTPUT_CHARS} characters"
                    break
            if rows == 0 and cap is None:
                out.write("(empty sheet)\n")
            if cap:
                notes.append(
                    f"{cap} in sheet '{ws.title}'; later rows and sheets are not included."
                )
                break
        return WorkbookText(text=out.getvalue().strip(), sheet_names=names, notes=notes)
    finally:
        wb.close()


# --- PPTX ----------------------------------------------------------------------

_SLIDE_RE = re.compile(r"^ppt/slides/slide(\d+)\.xml$")


def _paragraphs(xml: bytes) -> List[str]:
    root = ET.fromstring(xml)
    paragraphs = []
    for p in root.iter(f"{_A}p"):
        # Skip field runs (slide numbers, dates): <a:fld> wraps its own <a:t>.
        fields = {id(t) for fld in p.iter(f"{_A}fld") for t in fld.iter(f"{_A}t")}
        text = "".join(t.text or "" for t in p.iter(f"{_A}t") if id(t) not in fields)
        if text.strip():
            paragraphs.append(text.strip())
    return paragraphs


def _notes_member(zf: zipfile.ZipFile, slide_number: int) -> Optional[str]:
    rels = f"ppt/slides/_rels/slide{slide_number}.xml.rels"
    try:
        root = ET.fromstring(zf.read(rels))
    except KeyError:
        return None
    for rel in root.iter(f"{_REL}Relationship"):
        if (rel.get("Type") or "").endswith("/notesSlide"):
            return posixpath.normpath(
                posixpath.join("ppt/slides", rel.get("Target") or "")
            )
    return None


def pptx_text(data: bytes, include_notes: bool = True) -> str:
    zf = _open_zip(data)
    try:
        slides = sorted(
            (int(m.group(1)), name)
            for name in zf.namelist()
            if (m := _SLIDE_RE.match(name))
        )
        if not slides:
            raise ExtractionFailed("the presentation contains no slides.")
        blocks = []
        for number, name in slides:
            lines = [f"--- slide {number} ---", *_paragraphs(zf.read(name))]
            if include_notes:
                notes_name = _notes_member(zf, number)
                if notes_name and notes_name in zf.namelist():
                    notes = _paragraphs(zf.read(notes_name))
                    if notes:
                        lines.append("Speaker notes: " + " / ".join(notes))
            blocks.append("\n".join(lines))
            if sum(len(b) for b in blocks) > MAX_OUTPUT_CHARS:
                blocks.append(f"[output capped at {MAX_OUTPUT_CHARS} characters]")
                break
        return "\n\n".join(blocks)
    except (ET.ParseError, DefusedXmlException) as exc:
        raise ExtractionFailed("the presentation XML could not be parsed.") from exc
    finally:
        zf.close()
